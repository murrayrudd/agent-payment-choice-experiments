"""Create the Step 6 assessment using completed live results and reviewed interpretation."""
from pathlib import Path
from decimal import Decimal as D
from collections import Counter
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

B=Path(__file__).resolve().parent;R=B/'run_20261025'
S=json.loads((R/'summary.json').read_text());N=json.loads((B/'report_interpretation.json').read_text())
assert S['recorded_calls']==S['planned_calls'] and S['valid_responses']==S['recorded_calls']
d=Document();sec=d.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.75);sec.left_margin=sec.right_margin=Inches(.8);sec.footer_distance=Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Calibri');st._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Calibri')
d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.line_spacing=1.08;d.styles['Normal'].paragraph_format.space_after=Pt(6)
d.styles['Title'].font.size=Pt(19);d.styles['Title'].paragraph_format.space_after=Pt(6)
d.styles['Subtitle'].font.size=Pt(10.5);d.styles['Subtitle'].paragraph_format.space_after=Pt(12)
d.styles['Heading 1'].font.size=Pt(12.5);d.styles['Heading 1'].font.bold=True;d.styles['Heading 1'].paragraph_format.space_before=Pt(11);d.styles['Heading 1'].paragraph_format.space_after=Pt(5)
d.core_properties.title='Sol payment choice calibration pilot';d.core_properties.author='Murray A. Rudd';d.core_properties.subject='Step 6 of the Agents CE study rebuild'
def p(text,bold=False):
    x=d.add_paragraph();x.add_run(text).bold=bold;return x
def h(text,new=False):
    x=d.add_heading(text,1)
    if new:x.paragraph_format.page_break_before=True
def table(headers,rows,widths,small=10.5):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
    pr.append(borders);marg=OxmlElement('w:tblCellMar')
    for edge in ['top','bottom','left','right']:
        x=OxmlElement('w:'+edge);x.set(qn('w:w'),'80');x.set(qn('w:type'),'dxa');marg.append(x)
    pr.append(marg)
    for j,x in enumerate(headers):t.rows[0].cells[j].text=str(x)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for vals in rows:
        cells=t.add_row().cells
        for j,x in enumerate(vals):cells[j].text=str(x)
    for i,row in enumerate(t.rows):
        tr=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(tr)
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i==0:
                x=OxmlElement('w:shd');x.set(qn('w:fill'),'E8EEF3');c._tc.get_or_add_tcPr().append(x)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.03
                if j>0 and len(str(c.text))<22:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:run.font.size=Pt(small);run.bold=i==0
    d.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

d.add_paragraph('Sol payment choice calibration pilot','Title')
d.add_paragraph('Step 6 of the study rebuild  |  Version 0.1  |  October 25 2026','Subtitle')
p(N['opening'],True)
h('Main findings')
for x in N['main_findings']:p(x)
h('Collection and evidence')
table(['Measure','Result'],[
 ('Configuration','Sol 6.1  September 29 2026 snapshot'),
 ('Initial pilot calls',536),('Fixed presentation confirmation calls',96),('Targeted midpoint calls',len(S['adaptive_results'])*2),
 ('Valid responses',f"{S['valid_responses']}/{S['recorded_calls']}"),
 ('Total billed cost USD',f"${D(S['accounted_cost_usd']):.6f}"),
 ('Software spending ceiling','$2.00 for the whole stage'),
 ('Held-out prediction cases','32 cases from four offer pairs')
],[2.25,4.65])
p('The pilot retains instrument version 0.2. Qwen remains preliminary comparison evidence from Step 5.')
h('Decision for the next stage')
p(N['next_stage'])

h('Pilot design and analysis',True)
p('Every call used a fresh conversation. The common task, $1.00 principal, $1.30 initial-debit ceiling, certain conditional refund, 60-second refund delay, and defined private payment record stayed fixed. All offers were affordable. No numerical completion benefit, delay damage, or privacy harm was supplied.')
table(['Module','Calls','Purpose'],[
 ('Multidimensional choices',160,'20 offer pairs crossed with four mandates and two targets; 16 calibration and four held-out pairs.'),
 ('Exact core repetitions',32,'Repeat four calibration pairs within each mandate and target.'),
 ('Position and label contrasts',32,'Two label swaps and two position reversals within each mandate and target.'),
 ('Improvement fee ladders',280,'Six improvements; timeliness at both targets. Five fees, four mandates, two exact repetitions.'),
 ('Outside-option anchors',32,'Identical offers at favorable or unfavorable corners, both targets, four mandates, two repetitions.'),
 ('Adaptive midpoint checks',len(S['adaptive_results'])*2,'Two fresh responses at each eligible switching midpoint under the rule fixed before collection.'),
 ('Presentation confirmation',96,'Hold labels and positions fixed across mandates and targets for 16 calibration pairs; reuse 32 matching initial calls.')
],[1.55,.65,4.7])
p('The calibration pairs cover every proposed level and were selected for coverage and information under zero priors. They exclude simple attribute dominance. Four separate pairs were reserved before selection for prediction checks. The design is a pilot, not an optimized or powered main-study design.')
p('An audit during collection found that the initial balancing scheme changed presentation within some matched comparisons. A separate confirmation block was specified before outcome analysis. It holds presentation fixed within each pair. The original observations and primary prediction analysis remain unchanged; the additional block is not part of the original protocol.')
h('What the fee ladders identify')
p('The reference fee is $0.001. The improved offer has a fee of $0.001, $0.005, $0.020, $0.080, or $0.250, giving premiums from zero to $0.249. Each ladder changes one nonfee attribute, with all remaining terms fixed. Refund improvements are tested separately at 95% and 99% success. Timeliness is tested against 5- and 60-second targets.')
p('A series qualifies for one midpoint refinement only when both repetitions agree at every level, neither declines, the improvement wins at equal fees, and one monotone switch occurs. The rule permits at most 24 midpoints. A disagreement or reversal remains a result; it does not trigger response repair or range expansion.')
h('Exploratory prediction model')
p('Ridge multinomial logit uses only original first-repetition calibration choices. Each mandate has separate coefficients for purchase, fee, success, log time, returned principal, expected refund, log retention, and missing the target. Penalties are 0.1, 1, and 10; held-out choices do not select among them.')

h('Fee support and switching',True)
p('Cells report the observed premium bracket after any consistent midpoint refinement. A bracket gives an improvement choice at its lower end and a reference choice at its upper end. It is conditional on this contract, mandate, and presentation. It is not an independently validated operator WTP estimate.')
order=[('reliability_wide',60,'Success 95 to 99.9%'),('reliability_fine',60,'Success 99 to 99.9%'),('timeliness',5,'Time 30 to 5 s  target 5 s'),('timeliness',60,'Time 30 to 5 s  target 60 s'),('refund_p95',60,'Refund $0 to $1  success 95%'),('refund_p99',60,'Refund $0 to $1  success 99%'),('retention',60,'Retention 365 to 1 day')]
mandates=['Balanced','Economy','Continuity','Data minimization']
def laddercell(x):
    if x['category']=='monotone_switch':
        lo,hi=map(D,x['initial_bracket_premium_usd']);suffix=''
        a=next((a for a in S['adaptive_results'] if (a['ladder'],a['target_seconds'],a['mandate'])==(x['ladder'],x['target_seconds'],x['mandate'])),None)
        if a:
            if a['choices']==['beta','beta']:lo=D(a['premium'])
            elif a['choices']==['alpha','alpha']:hi=D(a['premium'])
            else:suffix=' *'
        return f'{lo:.3f} to {hi:.3f}'+suffix
    return {'improvement_at_maximum_premium':'Selected at 0.249','reference_at_equal_fee':'Reference at zero','repeat_disagreement_or_incomplete':'Repeat variation','nonmonotone':'Reversal','includes_decline':'Includes decline'}[x['category']]
table(['Improvement','Balanced','Economy','Continuity','Data minimization'],[(label,*[laddercell(next(x for x in S['ladders'] if (x['ladder'],x['target_seconds'],x['mandate'])==(name,target,m))) for m in mandates]) for name,target,label in order],[2.05,1.2,1.2,1.2,1.25],small=10)
p('All numbers are USD premiums above the reference fee. “Selected at 0.249” means the improvement remains selected at the largest tested premium. “Reference at zero” means the reference offer is selected at equal fees. “Repeat variation” means at least one identical-prompt pair disagrees.')
for x in N['fee_findings']:p(x)
h('Implications for attribute levels')
for x in N['level_decisions']:p(x)

h('Mandates targets and prediction',True)
p('The confirmation block holds physical offers, labels, and column order fixed across mandates and targets. These counts use 16 calibration pairs, giving 32 choices per mandate. They combine 96 additional calls with 32 matching initial calls. Held-out pairs are excluded from this block.')
cells=[]
C=S['fixed_presentation_confirmation'];assert C is not None
for m in mandates:
    purchased=sum(x['choice']!='DECLINE' for x in C['cells'] if x['mandate']==m);declined=sum(x['choice']=='DECLINE' for x in C['cells'] if x['mandate']==m)
    target=sum(x['changed'] for x in C['target_changes'] if x['mandate']==m)
    mandate=sum(x['choices'].get(m)!=x['choices'].get('Balanced') for x in C['mandate_changes'])
    cells.append((m,f'{purchased}/32',f'{declined}/32',f'{target}/16',f'{mandate}/32' if m!='Balanced' else 'Reference'))
table(['Mandate','Purchase','Decline','Target changes','Differs from Balanced'],cells,[1.8,.9,.8,1.5,1.9])
for x in N['mandate_findings']:p(x)
h('Withheld choice prediction')
table(['Penalty','Correct choices','Mean log loss','Majority baseline'],[(f"{x['penalty']:g}",f"{x['correct']}/{x['n']}",f"{x['mean_log_loss']:.3f}",f"{x['baseline_correct']}/{x['n']}") for x in S['prediction_summary']],[1,1.7,1.7,2.5])
p('Lower log loss is better. The baseline predicts the most common physical choice in each mandate and target among its calibration cases. It uses Laplace smoothing for probability scoring. The 32 observations come from only four held-out offer pairs; they do not establish broad external predictive validity.')
for x in N['prediction_findings']:p(x)
h('Scope of monetary interpretation')
p('Fee switches describe task-specific sacrifices between attributes. They do not validate operator preferences. Persistent endpoint choices cannot distinguish an absolute priority from a threshold beyond the tested range. The unstable fee estimates preclude defensible monetary coefficient ratios at this stage.')

h('Consistency outside option and next decisions',True)
cons=[]
for k,label in [('core','Exact core repetitions'),('ladder','Exact ladder repetitions'),('outside_anchor','Anchor purchase versus decline'),('adaptive_ladder','Exact midpoint repetitions')]:
    if k in S['repeat_consistency']:
        x=S['repeat_consistency'][k];cons.append((label,x['disagreements'],x['valid_pairs']))
for k,label in [('labels','Label swaps'),('position','Position reversals')]:
    x=S['presentation'][k];cons.append((label,x['disagreements'],x['valid_pairs']))
table(['Comparison','Disagreements','Matched pairs'],cons,[3.5,1.5,1.9])
p('Comparisons reuse cases and are not independent samples. Anchor comparisons collapse the two identical offers to PURCHASE, so arbitrary A/B tie-breaking does not count as a purchase inconsistency. Presentation contrasts are interpreted alongside ordinary repeat disagreement.')
for x in N['consistency_findings']:p(x)
h('Use of the outside option')
for x in N['outside_findings']:p(x)
h('Main study decision')
for x in N['final_decisions']:p(x)
h('Execution and replication')
p(f"The OpenAI Sol snapshot retained medium reasoning, an 8,192-token ceiling, provider-default sampling, explicit cache settings, and strict JSON output. Up to four requests ran concurrently. Billed cost of ${D(S['accounted_cost_usd']):.6f} reconciles exactly. The archive contains frozen inputs, exact requests and responses, separate confirmation and midpoint plans, scores, and offline analysis.")
p('The original protocol preceded collection; the confirmation plan preceded outcome analysis. All pilot observations remain separate from the future main study. Operator review, market calibration, and the proposed identity and linkability modules remain outstanding.')
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
for root in [d._element,d.styles._element]:
    for border in list(root.xpath('.//w:pBdr')):border.getparent().remove(border)
out=B.parent/'outputs/Agents_CE_Step_6_Sol_Calibration_Pilot_v01.docx';d.save(out);print(out)
