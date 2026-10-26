"""Create the main-run results report from observed data and reviewed interpretation."""
from pathlib import Path
from decimal import Decimal as D
from collections import Counter
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

B=Path(__file__).resolve().parents[1];R=B/'run_20261025'
S=json.loads((B/'assessment/report_metrics.json').read_text());N=json.loads((B/'assessment/interpretation.json').read_text());C=json.loads((B/'assessment/final_billing.json').read_text())
d=Document();sec=d.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.75);sec.left_margin=sec.right_margin=Inches(.8);sec.footer_distance=Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Calibri');st._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Calibri')
d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.line_spacing=1.04;d.styles['Normal'].paragraph_format.space_after=Pt(5)
d.styles['Title'].font.size=Pt(19);d.styles['Title'].paragraph_format.space_after=Pt(6)
d.styles['Subtitle'].font.size=Pt(10.5);d.styles['Subtitle'].paragraph_format.space_after=Pt(12)
d.styles['Heading 1'].font.size=Pt(12.5);d.styles['Heading 1'].font.bold=True;d.styles['Heading 1'].paragraph_format.space_before=Pt(11);d.styles['Heading 1'].paragraph_format.space_after=Pt(5)
d.core_properties.title='Agent payment choice main-run results';d.core_properties.author='Murray A. Rudd';d.core_properties.subject='Step 8 of the Agents CE study rebuild'
def p(text,bold=False):
    x=d.add_paragraph();x.add_run(text).bold=bold;return x
def h(text,new=False):
    x=d.add_heading(text,1)
    if new:x.paragraph_format.page_break_before=True
def table(headers,rows,widths,small=10.5):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
    pr.append(borders);marg=OxmlElement('w:tblCellMar')
    for edge in ['top','bottom','left','right']:
        x=OxmlElement('w:'+edge);x.set(qn('w:w'),'60');x.set(qn('w:type'),'dxa');marg.append(x)
    pr.append(marg)
    for j,x in enumerate(headers):t.rows[0].cells[j].text=str(x)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for vals in rows:
        cells=t.add_row().cells
        for j,x in enumerate(vals):cells[j].text=str(x)
    for i,row in enumerate(t.rows):
        tr=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(tr)
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i==0:
                x=OxmlElement('w:shd');x.set(qn('w:fill'),'E8EEF3');c._tc.get_or_add_tcPr().append(x)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.03
                if j>0 and len(str(c.text))<22:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:run.font.size=Pt(small);run.bold=i==0
    spacer=d.add_paragraph();spacer.paragraph_format.space_after=Pt(0);spacer.paragraph_format.line_spacing=Pt(2)
    return t

def pct(a,b):return f'{a/b:.1%}' if b else 'Not available'
def yn(x):return 'Yes' if x else 'No'
from datetime import datetime
from zoneinfo import ZoneInfo
def local(iso):return datetime.fromisoformat(iso.replace('Z','+00:00')).astimezone(ZoneInfo('America/Vancouver')).strftime('%B %d %Y at %I:%M %p')
d.add_paragraph('Agent payment choice main run results','Title')
d.add_paragraph('Step 8 of the study rebuild  Version 0.1  October 25 2026','Subtitle')
p(N['opening'],True)
for text in N['main_findings']:p(text)
h('Collection and evidence')
table(['Measure','Observed result'],[
 ('Model','Sol 6.1 September 29 2026 snapshot'),
 ('Completed and valid responses',f"{S['completed_calls']} completed; {S['valid_responses']} valid"),
 ('Estimation responses','512 first responses from 64 physical offer pairs'),
 ('Withheld prediction responses','256 first responses from 32 new physical offer pairs'),
 ('Repetitions and diagnostics','512 separate responses'),
 ('Declines across the complete run',f"{S['declines_all']} of {S['valid_responses']} valid responses"),
 ('Actual cost and spending ceiling',f"${D(S['total_cost_usd']):.6f}, reconciled exactly; $3.50 ceiling"),
 ('Account credit after collection',f"${D(C['available_credit_usd']):.6f}"),
 ('Collection period in Vancouver time',local(S['started_utc'])+' to '+local(S['ended_utc']))
],[2.45,4.45])
p('The run follows the frozen Step 7 design. Each call used a fresh conversation, one of four mandates, and either a 5-second or 60-second completion target. All offers remained affordable. Returned principal retained its positive wording. No prompt adaptation, automatic retry, or response replacement entered the run.')
h('Implication for the paper')
for text in N['paper_implications']:p(text)

h('Mandates and completion targets',True)
p('Comparisons hold offers, labels, and column order fixed across the 96 first-response core pairs. Counts describe the designed tasks, with no claim about choice frequencies in a market or operator population.')
table(['Mandate versus Balanced','5 second target','60 second target','Both targets'],[
 (x['mandate'],f"{x['by_target'][0]['changes']}/{x['by_target'][0]['pairs']}",f"{x['by_target'][1]['changes']}/{x['by_target'][1]['pairs']}",f"{x['changes']}/{x['paired_choices']} ({pct(x['changes'],x['paired_choices'])})") for x in S['mandate_changes']
],[2.4,1.35,1.35,1.8])
p('Entries count changes in the selected physical offer, including a change to or from DECLINE if one occurs. The same physical pairs contribute at both targets, so the 192 paired choices for a mandate are not 192 independent tasks.')
table(['Mandate','Changes when target shifts from 5 to 60 seconds','Share'],[
 (x['mandate'],f"{x['changes']}/{x['pairs']}",pct(x['changes'],x['pairs'])) for x in S['target_changes']
],[2.1,3.65,1.15])
for text in N['matched_interpretation']:p(text)
h('Attributes of selected offers')
table(['Mandate','Target seconds','Mean fee USD','Mean success percent','Mean retention days'],[
 (x['mandate'],x['target_seconds'],f"{x['mean_selected_attributes']['fee']:.3f}",f"{x['mean_selected_attributes']['success_probability']*100:.2f}",f"{x['mean_selected_attributes']['retention_days']:.1f}") for x in S['per_cell'] if x['mean_selected_attributes']
],[1.75,1.05,1.2,1.55,1.35],small=10)
p('These descriptive means combine correlated contract attributes. They do not isolate marginal effects. The archive retains exact selected-attribute changes and denominators.')

h('Prediction on withheld offer pairs',True)
p('Neither withheld answers nor diagnostic repetitions selected a penalty or entered estimation. Both prespecified fee specifications are reported. Training-only baselines provide direct comparisons. Accuracy uses the largest predicted probability; lower log loss indicates better probability predictions.')
names={'piecewise_fee':'Three-segment fee model','linear_fee':'Linear fee model','training_choice_share':'Training choice shares','lowest_fee':'Lowest fee rule','lowest_expected_expenditure':'Lowest expected expenditure rule','highest_success':'Highest success rule','meet_target_then_success':'Target then success and time rule','shortest_retention':'Shortest retention rule'}
table(['Model or rule','Correct of 256','Accuracy','Mean log loss'],[
 (names[x['model']],x['correct'],pct(x['correct'],x['n']),f"{x['mean_log_loss']:.3f}") for x in S['prediction_comparison']
],[3.1,1.15,1.2,1.45])
primary=next(x for x in S['prediction_scores'] if x['model']=='piecewise_fee');linear=next(x for x in S['prediction_scores'] if x['model']=='linear_fee');share=next(x for x in S['prediction_scores'] if x['model']=='training_choice_share')
table(['Mandate','Target seconds','Primary accuracy','Primary log loss','Share baseline log loss'],[
 (x['mandate'],x['target_seconds'],pct(round(x['accuracy']*x['n']),x['n']),f"{x['log_loss']:.3f}",f"{next(y['log_loss'] for y in share['by_cell'] if (y['mandate'],y['target_seconds'])==(x['mandate'],x['target_seconds'])):.3f}") for x in primary['by_cell']
],[1.75,1.05,1.35,1.3,1.45],small=10)
for text in N['prediction_interpretation']:p(text)
p('All 256 withheld responses purchase, so these results do not test prediction of declines. Archive intervals resample the 32 physical pairs with their cell observations retained and fitted predictions held fixed. They describe the finite task set, not human operators, other model snapshots, or different transaction environments.')

h('Repeatability and diagnostic contrasts',True)
labels={'core_estimation':'Exact estimation repeats','core_heldout':'Exact withheld repeats','fine_reliability':'Four-call fine success sets','outside_anchor':'Anchor purchase versus decline','refund_success':'Refund probe repeats'}
rep_rows=[(labels[x['module']],x['disagreeing_sets'],x['valid_sets'],str(x['calls_per_set'][0])) for x in S['repetition']]
rep_rows += [('Label swaps',x['changes'],x['valid_pairs'],'2') if x['variant']=='labels' else ('Column-order reversals',x['changes'],x['valid_pairs'],'2') for x in S['presentation']]
table(['Comparison','Disagreeing sets','Valid sets','Calls per set'],rep_rows,[3.1,1.35,1.2,1.25])
p('Sets disagree when valid choices differ. Four-call sets apply stricter unanimity than two-call sets. Identical anchor offers both count as PURCHASE. Interpret presentation effects alongside ordinary repeat variation.')
h('Fine success contrast')
def bracket(x):
    if x['premium_bracket_usd']:return '$'+x['premium_bracket_usd'][0]+' to $'+x['premium_bracket_usd'][1]
    return {'repeat_variation_or_incomplete':'Repeat variation','includes_decline':'Includes decline','reference_selected_at_all_fees':'Reference at all fees','higher_success_selected_at_maximum_fee':'Higher success at $0.249','nonmonotone':'Nonmonotone'}[x['category']]
table(['Mandate','5 second target','60 second target'],[
 (m,bracket(next(x for x in S['fine_reliability'] if x['mandate']==m and x['target_seconds']==5)),bracket(next(x for x in S['fine_reliability'] if x['mandate']==m and x['target_seconds']==60))) for m in ['Balanced','Economy','Continuity','Data minimization']
],[2.1,2.4,2.4],small=10)
p('Success rises from 99% to 99.9%, with a $0.50 return under both offers. This reduces expected refund by $0.0045. Brackets require unanimity in all four calls at each fee and one monotone switch. Premiums are above the $0.001 reference fee.')
h('Principal return probes')
p('Each contrast increases principal returned after failure from $0 to $1.')
refund_rows=[]
for pid in ['R1','R2','R3','R4']:
    rs=[x for x in S['refund_probes'] if x['pair_id']==pid];x=rs[0];choices=[c for r in rs for c in r['choices']]
    refund_rows.append((f"{D(x['success_probability'])*100:g}%",x['premium_usd'],x['net_expected_expenditure_change_usd'],f"{choices.count('beta')}/{len(choices)}"))
table(['Success','Premium USD','Change in expected expenditure USD','Full return selected'],refund_rows,[1.1,1.35,2.75,1.7],small=10)
for text in N['diagnostic_interpretation']:p(text)

h('Monetary interpretation and study limits',True)
p('The table applies the observable preliminary checks for monetary interpretation to the primary three-segment fee specification. A “Yes” for negative fee sensitivity requires every relevant fee segment to remain negative at all three prespecified penalties. The purchase-conditioned fit removes the unobserved outside-option margin.')
table(['Mandate','Target seconds','Full model fees negative','Conditional fees negative','Beats share baseline'],[
 (x['mandate'],x['target_seconds'],yn(x['all_full_fee_slopes_negative']),yn(x['all_conditional_fee_slopes_negative']),yn(x['beats_share_baseline'])) for x in S['fee_gates'] if x['specification']=='piecewise_fee'
],[1.75,1.05,1.4,1.4,1.3],small=10)
for text in N['monetary_interpretation']:p(text)
h('Outside option')
for text in N['outside_interpretation']:p(text)
h('What the evidence supports')
for text in N['scope_and_next']:p(text)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER;fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
for root in [d._element,d.styles._element]:
    for border in list(root.xpath('.//w:pBdr')):border.getparent().remove(border)
import sys
out=Path(sys.argv[1]) if len(sys.argv)>1 else B.parent/'outputs/Agents_CE_Step_8_Main_Run_Results_v01.docx'
out.parent.mkdir(parents=True,exist_ok=True);d.save(out);print(out)

