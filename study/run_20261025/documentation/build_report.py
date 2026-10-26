"""Create the Step 7 design report using frozen inputs and offline audits."""
from pathlib import Path
from decimal import Decimal as D
from collections import Counter
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

B=Path(__file__).resolve().parents[1]
P=json.loads((B/'inputs/protocol.json').read_text());A=json.loads((B/'development/design_audit.json').read_text());S=json.loads((B/'development/simulation_results.json').read_text());C=json.loads((B/'development/budget_check.json').read_text())
d=Document();sec=d.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.75);sec.left_margin=sec.right_margin=Inches(.8);sec.footer_distance=Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Calibri');st._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Calibri')
d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.line_spacing=1.04;d.styles['Normal'].paragraph_format.space_after=Pt(5)
d.styles['Title'].font.size=Pt(19);d.styles['Title'].paragraph_format.space_after=Pt(6)
d.styles['Subtitle'].font.size=Pt(10.5);d.styles['Subtitle'].paragraph_format.space_after=Pt(12)
d.styles['Heading 1'].font.size=Pt(12.5);d.styles['Heading 1'].font.bold=True;d.styles['Heading 1'].paragraph_format.space_before=Pt(11);d.styles['Heading 1'].paragraph_format.space_after=Pt(5)
d.core_properties.title='Agent payment choice main-study design';d.core_properties.author='Murray A. Rudd';d.core_properties.subject='Step 7 of the Agents CE study rebuild'
def p(text,bold=False):
    x=d.add_paragraph();x.add_run(text).bold=bold;return x
def h(text,new=False):
    x=d.add_heading(text,1)
    if new:x.paragraph_format.page_break_before=True
def table(headers,rows,widths,small=10.5):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
    pr.append(borders);marg=OxmlElement('w:tblCellMar')
    for edge in ['top','bottom','left','right']:
        x=OxmlElement('w:'+edge);x.set(qn('w:w'),'60');x.set(qn('w:type'),'dxa');marg.append(x)
    pr.append(marg)
    for j,x in enumerate(headers):t.rows[0].cells[j].text=str(x)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for vals in rows:
        cells=t.add_row().cells
        for j,x in enumerate(vals):cells[j].text=str(x)
    for i,row in enumerate(t.rows):
        tr=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(tr)
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i==0:
                x=OxmlElement('w:shd');x.set(qn('w:fill'),'E8EEF3');c._tc.get_or_add_tcPr().append(x)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.03
                if j>0 and len(str(c.text))<22:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:run.font.size=Pt(small);run.bold=i==0
    spacer=d.add_paragraph();spacer.paragraph_format.space_after=Pt(0);spacer.paragraph_format.line_spacing=Pt(2)
    return t

d.add_paragraph('Agent payment choice main-study design','Title')
d.add_paragraph('Step 7 of the study rebuild  |  Version 0.1  |  October 25 2026','Subtitle')
p('Freeze a 1,280-call main study using Sol 6.1 alone: 64 estimation pairs, 32 new pairs reserved for prediction, and targeted repeatability and interpretation checks.',True)
p('This stage prepares the design and executable package. No main-study model responses have been collected, and no API credit was spent on Step 7.')
h('Research question and role')
p(P['research_question'],True)
p('The research question remains suitable. The main contribution is evidence about attribute trade-offs and their dependence on an operator mandate. Completion targets establish a specific contextual boundary. Prediction tests assess whether the fitted account travels to new combinations; monetary interpretation remains conditional.')
table(['Supporting question','Role in the main study'],[
 ('Attribute effects and interactions','Estimate choices across all five attributes, including the interaction between success and refunded principal.'),
 ('Mandate-dependent trade-offs','Compare the same physical offers under Balanced, Economy, Continuity, and Data minimization mandates.'),
 ('Transaction-condition boundaries','Compare a 5-second with a 60-second desired completion target while keeping affordability and other terms fixed.'),
 ('Prediction and monetary interpretation','Test new offer combinations. Consider task-specific fee compensation only if price sensitivity and predictive checks support it.')
],[2.2,4.7])
h('What changes after the pilot')
p('Retain the positive “principal returned after failure” wording and all five attributes. Add four interior fee levels to resolve switches within the existing $0.001–$0.250 range. Keep labels and column order fixed within every matched mandate and target comparison. Use entirely new core offer pairs for estimation and withheld prediction.')
p('The pilot yielded 672 valid choices, cost $1.387236, and produced no declines. Its fixed-presentation confirmation showed mandate and target differences, while fine success contrasts produced some repeat variation. Those findings motivate the design; they are not main-study observations.')
h('Budget decision')
p('The expected main-run cost is $2.64 at the observed pilot mean, or $3.30 with a 25% allowance. The software ceiling is $3.50. Available credit was $3.90 when checked on October 25 at 1:03 p.m. Vancouver time. A top-up is not needed for the planned run under these estimates; collection must stop if its ceiling cannot reserve the next calls.')

h('Instrument and matched conditions',True)
table(['Attribute','Frozen main-study levels'],[
 ('Payment service fee, USD','$0.001, $0.003, $0.005, $0.013, $0.020, $0.050, $0.080, $0.165, $0.250'),
 ('Probability of usable data','95%, 99%, 99.9%; failures per 1,000 also shown'),
 ('Completion time','1, 5, 30 seconds'),
 ('Principal returned after failure','$0, $0.50, $1.00'),
 ('Payment-record retention','1, 30, 365 days')
],[2.35,4.55])
p('The $1.00 service principal and $1.30 maximum initial debit stay fixed. All offers are affordable. The service, data content, quality, rights, and feasibility are identical. Each call permits one purchase or DECLINE, with no second attempt or substitute. Declining leaves the report incomplete but has no stated numerical penalty.')
p('The payment fee is nonrefundable. Principal return is certain conditional on failure and occurs automatically 60 seconds after the definitive failure notice. Successful purchases receive no refund. The narrowly defined private payment record and durable operator receipt are unchanged from instrument version 0.2.')
h('Mandates and completion targets')
table(['Mandate','Priority, with other attributes still considered'],[
 ('Balanced','No single attribute receives an additional priority.'),
 ('Economy','Give greater weight to limiting operator expenditure.'),
 ('Continuity','Give greater weight to dependable, timely task completion.'),
 ('Data minimization','Give greater weight to limiting retention and use of payment information.')
],[1.55,5.35])
p('The package preserves the exact mandate wording. The desired completion target is either 5 or 60 seconds; later data remain useful. Mandates express preferences, not additional feasibility restrictions. Every core pair appears in all eight mandate–target cells.')
h('Interpret returned principal jointly with success')
p('Expected expenditure equals the principal plus fee minus the expected principal return: 1 + fee − (1 − success probability) × returned principal. At a $0.50 return, raising success from 95% to 99.9% increases expected expenditure by $0.0245 at equal fees because failure refunds become less frequent. Raising success from 99% to 99.9% changes it by $0.0045.')
p('A lower-success choice under an expenditure-focused mandate is therefore not, by itself, evidence of misunderstanding. The design retains success and returned principal separately and includes expected refund as an interaction term. No task-completion value is invented to force switches.')

h('Call allocation and design construction',True)
table(['Module','Calls','Allocation'],[
 ('Estimation choices',512,'64 new pairs × 8 cells'),
 ('Withheld prediction choices',256,'32 new pairs × 8 cells'),
 ('Exact estimation repetitions',128,'16 selected pairs × 8 cells'),
 ('Exact withheld repetitions',64,'8 selected pairs × 8 cells'),
 ('Label swaps',32,'4 estimation pairs × 8 cells'),
 ('Column-order reversals',32,'4 other estimation pairs × 8 cells'),
 ('Fine success contrast',160,'5 fees × 8 cells × 4 exact repetitions'),
 ('Outside-option anchors',32,'2 corners × 8 cells × 2 repetitions'),
 ('Refund–success probes',64,'4 contrasts × 8 cells × 2 repetitions'),
 ('Total',1280,'One fixed Sol configuration')
],[2.3,.65,3.95])
p('The 64 withheld repeat responses assess repeatability; they are not additional prediction cases. Diagnostic calls never enter core fitting or withheld scoring. Identical-offer anchors are evaluated as purchase versus decline, so arbitrary A/B tie-breaking does not imply inconsistency.')
h('New pairs and presentation control')
p('The profile grid contains 729 contracts. Candidate pairs differ on at least two attributes and exclude all 78 distinct physical pairs encountered in Steps 5–6. Thirty-two withheld pairs are selected by a coverage-constrained random draw before information optimization. Individual profiles may recur; withheld prediction concerns new combinations within the same attribute levels.')
p('Within each core split, the four combinations of label assignment and column order are equally represented. Each pair keeps its presentation across all eight cells and exact repetitions. Label and order diagnostics change one presentation factor at a time. All 1,280 fresh-context calls receive a fixed randomized order.')
h('Information criterion')
p('Greedy selection uses information averaged equally over a zero prior and all three pilot coefficient sets, with penalties 0.1, 1, and 10. No pilot held-out result selects a preferred prior. Coverage requirements protect interior fee levels and all nonfee levels. The criterion is a transparent heuristic; it does not establish global optimality.')
p('The conditional design has 12 terms: three fee segments with knots at $0.02 and $0.08, two terms for each three-level nonfee attribute, and expected refund. The planned multinomial model adds a purchase constant. Separate cell coefficients allow mandate and target interactions with every attribute.')

h('Offline checks and sample-size decision',True)
table(['Check','Result'],[
 ('Fresh core pairs','96 of 96; zero overlap with development pairs'),
 ('Attribute coverage','All nine fees and all nonfee levels in both core splits'),
 ('Conditional feature rank','12 of 12 at 48, 64, and 80 estimation pairs'),
 ('Information condition numbers','37.3–38.8 across the eight cells'),
 ('Relative information score','1.62 times the geometric information score of 100 random 64-pair designs, on average'),
 ('Random comparisons outperforming selected design','0 of 100 under the same candidate pool and assumed criterion')
],[3.2,3.7])
p('These are design diagnostics under stated assumptions, not observed behavioral results. The relative score uses the twelfth root of the information determinant and includes the same small numerical stabilizer for every compared design.')
h('Synthetic comparison of 48, 64, and 80 estimation pairs')
p('Forty synthetic replications per cell compare smooth pilot-like choices, sharper pilot-like choices, and mandate-priority rules with a 5% random-choice lapse. Prediction is evaluated on 1,000 additional pairs excluded from both the optimized design and the main withheld set. Only purchase-conditioned choices are simulated; declines are not modeled in this exercise.')
rows=[]
for scenario,label in [('smooth_pilot','Smooth pilot utility'),('sharp_pilot','Sharper pilot utility'),('priority_with_lapse','Priority rule with lapse')]:
    values=[next(r for r in S['summary'] if r['scenario']==scenario and r['estimation_pairs']==n) for n in [48,64,80]]
    rows.append((label,*[f"{r['expected_log_loss']:.3f}" for r in values]))
table(['Assumed mechanism','48 pairs','64 pairs','80 pairs'],rows,[2.7,1.4,1.4,1.4])
p('Entries are mean expected log loss; lower is better. Moving from 48 to 64 pairs improves all three scenarios. Moving to 80 improves them further, with smaller gains. This supports 64 as a budget-conscious working allocation, not as a uniquely optimal or formally powered sample.')
table(['Estimation pairs','Total calls','Pilot-rate cost','With 25% allowance'],[
 (48,1152,'$2.38','$2.97'),(64,1280,'$2.64','$3.30'),(80,1408,'$2.91','$3.63')
],[1.6,1.3,1.8,2.2])
p('The comparison holds the withheld and diagnostic allocations fixed. The 80-pair option would exceed the $3.50 stage ceiling under the 25% cost allowance. The 64-pair design retains 32 separate prediction pairs and the diagnostic checks that the pilot showed were useful.')

h('Prespecified analysis and interpretation',True)
h('Matched choice patterns')
p('Use first scheduled core responses to compare physical choices under different mandates and targets. Report choice-change rates over the 96 pairs, including declines. Report changes in the attributes of selected offers only when both matched choices purchase, with that denominator. Interpret these patterns alongside ordinary repeat variation and separate label and order checks.')
h('Estimation and prediction')
p('Fit ridge multinomial logit separately in all eight cells, using only the 64 estimation pairs. The primary specification uses three continuous fee segments; a linear-fee model is a prespecified sensitivity. Both include categorical nonfee effects, expected refund, and a purchase constant. The outside option has utility zero. All coefficients are penalized.')
p('Four fixed folds each contain 16 physical estimation pairs, preserving their eight-cell membership. For each specification, select a common penalty from 0.1, 1, and 10 using mean validation-fold log loss. Save training decisions before scoring the 32 withheld pairs. Report both specifications; withheld results cannot select a new model or penalty.')
p('Evaluate accuracy and log loss across all three options, overall and by cell. Compare with smoothed training choice shares and five simple attribute rules: lowest fee, lowest expected expenditure, highest success, meeting the target before success and time, and shortest retention. Rule ties and training-only lapse calibration are fixed in the protocol.')
h('Failures, repetitions, and uncertainty')
p('Keep schema failures, refusals, truncation, and DECLINE distinct. Never replace an invalid first response with its exact repeat. Require at least 48 valid estimation responses in every cell to fit the models; report all missing denominators. Diagnostics and repeated calls do not increase the estimation or prediction sample.')
p('Resample physical pairs with their cell observations retained, using 1,000 fixed-seed draws for descriptive and prediction-score intervals. These describe sensitivity to the finite task set; they are not population confidence intervals for human operators. Prediction-score resampling holds the fitted model fixed. Raw utility coefficients across separately fitted cells do not establish differences in preference intensity.')
h('Monetary interpretation remains subordinate')
p('Report observed switching brackets, repeated-choice variation, reversals, and persistence at the tested endpoints. Do not automatically report WTP ratios. Task-specific fee compensation requires a unique interior crossing, stable negative fee gradients under the prespecified penalties and purchase-conditioned sensitivity, better withheld log loss than the training-share baseline in the relevant cell, and bounded within-range uncertainty.')
p('If there are no declines, the purchase constant and outside-choice probabilities remain regularization dependent. No reservation price or operator welfare estimate follows from universal purchase. Any failed monetary-interpretation gate means that the corresponding estimate is withheld.')

h('Execution record and next stage',True)
table(['Item','Frozen setting or current check'],[
 ('Model','openai/gpt-6.1-sol-20260929; OpenAI endpoint only'),
 ('Generation','Medium reasoning; 8,192 output-token ceiling; provider-default sampling; exact tested cache settings'),
 ('Response','Strict JSON choice A, B, or DECLINE; fresh conversation each time'),
 ('Concurrency','At most four calls; reserve every in-flight request'),
 ('Stage spending ceiling','$3.50; no automatic retries or budget expansion'),
 ('Available account credit','$3.899145 at October 25, 1:03 p.m. Vancouver time'),
 ('Expected cost / allowance','$2.642354 / $3.302943 with 25% allowance'),
 ('Key expiry','November 1, 2026, 9:33 a.m. Vancouver time'),
 ('Step 7 API model calls','0')
],[2.15,4.75])
p('A read-only check confirmed the snapshot and rates: $2 per million input tokens and $10 per million output tokens. Before collection, the runner rechecks the endpoint, prices, expiry, credit, and hashes. It stops scheduling on insufficient budget or transport, routing, billing, or response ambiguity.')
h('Frozen package and verification')
p('The archive contains the instrument, 1,280 prompts and requests, assignments, protocol, development inputs, audits, simulations, and runnable collection and analysis scripts. A SHA-256 manifest records the frozen files. Credentials and main-study responses are absent.')
p('Checks verified accounting, coverage, new pairs, matched presentation, exact repeats, requests, and call counts. Mock-provider tests exercised budget stops and preserved failures. Synthetic analysis confirmed that changing withheld and diagnostic answers leaves training fits unchanged, and that invalid first responses are not replaced by repeats.')
h('Scope and next action')
p('The next stage is the frozen main collection, followed by the prespecified analysis. It does not add Qwen or change the question after viewing outcomes. Findings will describe one model snapshot acting under supplied mandates in one defined task. Operator review, market-fee calibration, identity and linkability, and different refund delays remain separate extensions.')
h('Methodological basis')
p('Johnson FR et al. (2013), Value in Health 16:3–13, doi:10.1016/j.jval.2012.08.2223: design should support the intended choice model, with sensitivity to incorrect priors. These principles motivate feature coverage and prior sensitivity; they do not validate this particular design.')
p('Rose JM and Bliemer MCJ (2013), Transportation 40:1021–1041, doi:10.1007/s11116-013-9451-z: sample requirements depend on design assumptions. Access was limited to the publisher abstract and metadata. Human-panel calculations are not transferred to model calls. Source URLs are retained in protocol.json.')
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
for root in [d._element,d.styles._element]:
    for border in list(root.xpath('.//w:pBdr')):border.getparent().remove(border)
import sys
out=Path(sys.argv[1]) if len(sys.argv)>1 else B/'documentation/Agents_CE_Step_7_Main_Study_Design_v01.docx'
out.parent.mkdir(parents=True,exist_ok=True);d.save(out);print(out)

